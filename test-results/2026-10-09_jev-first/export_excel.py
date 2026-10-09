import json
from collections import Counter
from pathlib import Path
from statistics import mean, median
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
root=Path(__file__).parent
data=json.loads((root/'raw_results.json').read_text(encoding='utf-8-sig'))
rows=data['records']; lats=sorted(r['latencySeconds'] for r in rows if r['latencySeconds'] is not None)
p95=lats[min(len(lats)-1,round(.95*(len(lats)-1)))]
models=Counter((r.get('decisionModel') or 'unreported') for r in rows)
summary={'startedAtIst':data['startedAtIst'],'completedAtIst':data['completedAtIst'],'total':len(rows),'passed':sum(bool(r['passed']) for r in rows),'failed':sum(not bool(r['passed']) for r in rows),'failedCases':[r['caseId'] for r in rows if not r['passed']],'latency':{'min':min(lats),'average':round(mean(lats),4),'p50':round(median(lats),4),'p95':p95,'max':max(lats)},'tokens':{'input':sum(r.get('inputTokens') or 0 for r in rows),'output':sum(r.get('outputTokens') or 0 for r in rows),'total':sum(r.get('totalTokens') or 0 for r in rows)},'models':dict(models),'resourceSnapshots':data['resourceSnapshots']}
(root/'summary.json').write_text(json.dumps(summary,indent=2),encoding='utf-8')
wb=Workbook(); ws=wb.active;ws.title='Summary'
items=[('Test suite','chat_bot_5 Jev-first stateful stress test'),('Started IST',summary['startedAtIst']),('Completed IST',summary['completedAtIst']),('Primary model','jevos-v4 / OpenVINO INT8'),('Fallback model','qwen2.5:1.5b / local Ollama'),('Total turns',summary['total']),('Passed',summary['passed']),('Failed',summary['failed']),('Failed cases',', '.join(summary['failedCases'])),('Minimum response (s)',summary['latency']['min']),('Average response (s)',summary['latency']['average']),('P50 response (s)',summary['latency']['p50']),('P95 response (s)',summary['latency']['p95']),('Maximum response (s)',summary['latency']['max']),('Input tokens',summary['tokens']['input']),('Output tokens',summary['tokens']['output']),('Jev turns',models.get('jevos-v4/openvino-int8',0)),('Qwen fallback turns',models.get('qwen2.5:1.5b-fallback',0))]
for x in items:ws.append(x)
log=wb.create_sheet('Chat Log');heads=['Timestamp IST','Phase','Case','Query','Expected Asset','Expected Period','Expected Action','Passed','Response Time (s)','Model','Input Tokens','Output Tokens','Total Tokens','Actual Action','Reply','Parameters','Job ID','Error'];log.append(heads)
for r in rows:log.append([r['timestampIst'],r['phase'],r['caseId'],r['query'],r['expectedAsset'],r['expectedPeriod'],r['expectedAction'],r['passed'],r['latencySeconds'],r.get('decisionModel'),r.get('inputTokens'),r.get('outputTokens'),r.get('totalTokens'),r.get('suggestedAction'),r.get('reply'),json.dumps(r.get('parameters'),ensure_ascii=False),r.get('jobId'),r.get('error')])
res=wb.create_sheet('Resources');res.append(['Timestamp IST','Phase','Container','CPU','Memory Usage','Memory %','Network I/O','Block I/O','PIDs'])
for sample in data['resourceSnapshots']:
 for c in sample['containers']:res.append([sample['timestampIst'],sample['label'],c.get('Name') or c.get('Container'),c.get('CPUPerc'),c.get('MemUsage'),c.get('MemPerc'),c.get('NetIO'),c.get('BlockIO'),c.get('PIDs')])
for sh in wb.worksheets:
 sh.freeze_panes='A2';sh.auto_filter.ref=sh.dimensions
 for c in sh[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='1F4E78');c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
 for row in sh.iter_rows(min_row=2):
  for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
 for col in range(1,sh.max_column+1):sh.column_dimensions[get_column_letter(col)].width=min(max(len(str(sh.cell(r,col).value or '')) for r in range(1,sh.max_row+1))+2,48)
out=root/'chat_bot_5_jev_stress_test.xlsx';wb.save(out);print(json.dumps(summary,indent=2));print(out)
